"""Create the sample .docx templates in templates_docx/ (NDA + two Italian HR compliance documents)."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path(__file__).parent / "templates_docx"


def _doc():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    return d


def _sign_block(doc, left="{{ company_name }}", right="{{ full_name }}"):
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).text = left; t.cell(0, 1).text = "Il/La dipendente"
    t.cell(1, 0).text = "{{ company_signatory }}"; t.cell(1, 1).text = right
    t.cell(2, 0).text = "Firma: ______________________"; t.cell(2, 1).text = "Firma: ______________________"
    doc.add_paragraph("")
    doc.add_paragraph("Luogo e data: {{ sede_lavoro }}, {{ today_it }}")


def nda():
    doc = _doc()
    doc.add_heading("MUTUAL NON-DISCLOSURE AGREEMENT / ACCORDO DI RISERVATEZZA", level=1)
    doc.add_paragraph("Between / Tra")
    doc.add_paragraph("{{ company_name }}, with registered office at {{ company_address }}, VAT no. {{ company_vat }}, "
                      "represented by {{ company_signatory }} (\"Company\")")
    doc.add_paragraph("and / e")
    doc.add_paragraph("{{ counterparty_name }}, with registered office at {{ counterparty_address }}, VAT no. {{ counterparty_vat }}, "
                      "represented by {{ counterparty_signatory }} (\"Counterparty\")")
    doc.add_paragraph("Effective date / Data di efficacia: {{ effective_date | date }}")
    doc.add_heading("1. Purpose / Finalità", level=2)
    doc.add_paragraph("The Parties intend to exchange confidential information in connection with: {{ purpose }}.")
    doc.add_heading("2. Term / Durata", level=2)
    doc.add_paragraph("This Agreement remains in force for {{ term_years }} year(s) from the Effective Date. "
                      "Confidentiality obligations survive for {{ survival_years }} year(s) after termination.")
    doc.add_heading("3. Governing law / Legge applicabile", level=2)
    doc.add_paragraph("This Agreement is governed by the laws of {{ governing_law }}. Exclusive jurisdiction: Court of {{ jurisdiction }}.")
    doc.add_paragraph("{% if penalty_amount %}4. Penalty / Penale: in case of breach, the defaulting Party shall pay "
                      "{{ penalty_amount | money }}, without prejudice to further damages.{% endif %}")
    doc.add_paragraph("")
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "{{ company_name }}"; t.cell(0, 1).text = "{{ counterparty_name }}"
    t.cell(1, 0).text = "{{ company_signatory }}"; t.cell(1, 1).text = "{{ counterparty_signatory }}"
    t.cell(2, 0).text = "Signature: ______________________"; t.cell(2, 1).text = "Signature: ______________________"
    doc.add_paragraph("")
    doc.add_paragraph("Generated on {{ today_it }}").runs[0].font.size = Pt(8)
    doc.save(OUT / "nda_mutual.docx")


def informativa_privacy():
    doc = _doc()
    doc.add_heading("INFORMATIVA AI DIPENDENTI SUL TRATTAMENTO DEI DATI PERSONALI", level=1)
    doc.add_paragraph("ai sensi degli artt. 13 e 14 del Regolamento (UE) 2016/679 (\"GDPR\") e del D.Lgs. 196/2003 e s.m.i.")
    doc.add_paragraph("Titolare del trattamento: {{ company_name }}, con sede legale in {{ company_address }}, P.IVA {{ company_vat }}, "
                      "e-mail {{ privacy_email }}.")
    doc.add_paragraph("{% if dpo_email %}Responsabile della protezione dei dati (DPO): {{ dpo_email }}.{% endif %}")
    doc.add_paragraph("Interessato: {{ full_name }}, C.F. {{ codice_fiscale }}, {{ job_title }} – {{ department }}, "
                      "assunto/a il {{ hire_date | date }}, sede di lavoro {{ sede_lavoro }}.")
    doc.add_heading("1. Finalità e basi giuridiche", level=2)
    doc.add_paragraph("I dati sono trattati per la gestione del rapporto di lavoro (art. 6, par. 1, lett. b e c GDPR), per adempiere "
                      "a obblighi di legge in materia fiscale, previdenziale e di sicurezza sul lavoro (D.Lgs. 81/2008), e per il legittimo "
                      "interesse del Titolare alla sicurezza dei sistemi informativi (art. 6, par. 1, lett. f GDPR). Le categorie particolari "
                      "di dati (art. 9 GDPR) sono trattate nei limiti dell'art. 9, par. 2, lett. b GDPR e dell'art. 2-sexies D.Lgs. 196/2003.")
    doc.add_heading("2. Destinatari e trasferimenti", level=2)
    doc.add_paragraph("I dati possono essere comunicati a consulenti del lavoro, istituti previdenziali e assicurativi, medico competente, "
                      "fornitori di servizi IT nominati responsabili ex art. 28 GDPR, e alle società del gruppo {{ group_name }}. Eventuali "
                      "trasferimenti extra-UE avvengono sulla base delle clausole contrattuali standard (art. 46 GDPR).")
    doc.add_heading("3. Conservazione", level=2)
    doc.add_paragraph("I dati sono conservati per la durata del rapporto e, successivamente, per i termini di legge (di regola 10 anni dalla "
                      "cessazione per la documentazione contabile e previdenziale).")
    doc.add_heading("4. Diritti dell'interessato", level=2)
    doc.add_paragraph("L'interessato può esercitare i diritti di cui agli artt. 15-22 GDPR scrivendo a {{ privacy_email }} e proporre reclamo "
                      "al Garante per la protezione dei dati personali (art. 77 GDPR).")
    doc.add_paragraph("Versione informativa: {{ policy_version }}")
    doc.add_heading("Dichiarazione di ricezione", level=2)
    doc.add_paragraph("Il/La sottoscritto/a {{ full_name }} dichiara di aver ricevuto e letto la presente informativa.")
    _sign_block(doc)
    doc.save(OUT / "informativa_privacy_dipendenti.docx")


def consegna_policy():
    doc = _doc()
    doc.add_heading("RICEVUTA DI CONSEGNA E PRESA VISIONE DELLA DOCUMENTAZIONE AZIENDALE", level=1)
    doc.add_paragraph("Il/La sottoscritto/a {{ full_name }}, C.F. {{ codice_fiscale }}, in qualità di {{ job_title }} presso "
                      "{{ company_name }}, sede di {{ sede_lavoro }},")
    doc.add_paragraph("DICHIARA")
    doc.add_paragraph("di aver ricevuto in data odierna copia dei seguenti documenti, di averne preso visione e di impegnarsi a rispettarne il contenuto:")
    doc.add_paragraph("{{ documents_list | lines }}")
    doc.add_paragraph("Versione: {{ policy_version }}. I documenti sono inoltre disponibili su {{ intranet_url }}.")
    doc.add_paragraph("Il/La sottoscritto/a è consapevole che la violazione delle suddette disposizioni può costituire illecito disciplinare "
                      "ai sensi dell'art. 7 L. 300/1970 e del CCNL applicato ({{ ccnl }}), nonché, ove rilevante, dei presidi adottati ai sensi "
                      "del D.Lgs. 231/2001.")
    _sign_block(doc)
    doc.save(OUT / "consegna_policy.docx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    nda(); informativa_privacy(); consegna_policy()
    print("\n".join(str(p) for p in sorted(OUT.glob("*.docx"))))
