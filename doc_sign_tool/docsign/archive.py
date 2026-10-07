"""Archive signed copies and export the compliance register (备案登记表)."""
from __future__ import annotations

import csv
import re
import shutil
from datetime import datetime
from pathlib import Path

from .config import Config
from .store import Store


def _safe(s: str) -> str:
    return re.sub(r"[^\w\-. ]+", "_", s or "", flags=re.UNICODE).strip(" ._") or "x"


def archive_signed(cfg: Config, store: Store, rec: dict, src: Path, note: str | None = None) -> Path:
    """Copy a signed file into archive/<employee>/<title>_<date>_SIGNED.<ext> and mark the record signed."""
    if rec.get("employee_id"):
        folder = cfg.archive_dir / _safe(f"{rec.get('employee_name') or ''}_{rec['employee_id']}")
    else:
        folder = cfg.archive_dir / "_senza_dipendente"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d")
    dest = folder / f"{_safe(rec['title'])}_{stamp}_SIGNED{src.suffix.lower()}"
    i = 2
    while dest.exists():
        dest = folder / f"{_safe(rec['title'])}_{stamp}_SIGNED_{i}{src.suffix.lower()}"
        i += 1
    shutil.copy2(src, dest)
    store.mark_signed(rec["id"], str(src), note, archive_path=str(dest))
    return dest


REGISTER_COLUMNS = [
    ("record_id", "ID"), ("employee_id", "Matricola"), ("employee_name", "Dipendente"), ("email", "E-mail"),
    ("title", "Documento"), ("template", "Template"), ("status", "Stato"), ("created_at", "Generato"),
    ("sent_at", "Inviato"), ("signed_at", "Firmato"), ("archive_path", "File archiviato"), ("note", "Note"),
]


def _rows(store: Store) -> list[dict]:
    out = []
    for r in store.list(limit=100000):
        signer = next((x for x in r["recipients"] if x["role"] == "signer"), {})
        out.append({
            "record_id": r["id"], "employee_id": r["employee_id"] or "", "employee_name": r["employee_name"] or signer.get("name", ""),
            "email": signer.get("email", ""), "title": r["title"], "template": r["template"], "status": r["status"],
            "created_at": r["created_at"], "sent_at": r["sent_at"] or "", "signed_at": r["signed_at"] or "",
            "archive_path": r["archive_path"] or "", "note": r["note"] or "",
        })
    return out


def export_register(cfg: Config, store: Store, fmt: str = "xlsx") -> Path:
    cfg.archive_dir.mkdir(parents=True, exist_ok=True)
    rows = _rows(store)
    stamp = datetime.now().strftime("%Y-%m-%d")
    if fmt == "csv":
        dest = cfg.archive_dir / f"registro_firme_{stamp}.csv"
        with open(dest, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow([lbl for _, lbl in REGISTER_COLUMNS])
            for r in rows:
                w.writerow([r[k] for k, _ in REGISTER_COLUMNS])
        return dest
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro firme"
    ws.append([lbl for _, lbl in REGISTER_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
    for r in rows:
        ws.append([r[k] for k, _ in REGISTER_COLUMNS])
    for col, width in zip("ABCDEFGHIJKL", (24, 12, 28, 30, 36, 24, 10, 20, 20, 20, 60, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    # Matrix sheet: employee × document, latest status
    emps = store.list_employees(include_inactive=True)
    latest = store.latest_by_employee_template()
    templates = sorted({t for _, t in latest})
    m = wb.create_sheet("Matrice")
    m.append(["Matricola", "Dipendente", "E-mail"] + templates)
    for c in m[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
    for e in emps:
        m.append([e["employee_id"], e["full_name"], e["email"]] +
                 [(latest.get((e["employee_id"], t)) or {}).get("status", "—") for t in templates])
    m.freeze_panes = "D2"
    dest = cfg.archive_dir / f"registro_firme_{stamp}.xlsx"
    wb.save(dest)
    return dest
